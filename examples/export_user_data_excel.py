"""
Letterboxd User Data Excel Exporter

Exports comprehensive user data from Letterboxd profiles directly to an Excel (.xlsx) file.
- Multi-sheet workbook (Overview, Films, Diary, Reviews, Watchlist, Likes, Social)
- Formatted tables with auto-filters and frozen header rows
- Column width auto-sizing with text wrapping
- Fast export reusing cached JSON files if available
"""

__title__ = "User Data Excel Exporter"
__description__ = "Exports comprehensive user data from Letterboxd profiles to a styled Excel workbook."
__version__ = "0.1.0"
__author__ = "fastfingertips"
__author_url__ = "https://github.com/fastfingertips"
__created_at__ = "2026-10-19"

import argparse
import inspect
import os
import sys
import time
import webbrowser

from fastfingertips.terminal_utils import args_exists, get_input

from letterboxdpy import user
from letterboxdpy.utils.utils_directory import check_and_create_dirs
from letterboxdpy.utils.utils_file import JsonFile, build_click_url, build_path
from letterboxdpy.utils.utils_string import strip_prefix

try:
    import openpyxl
    from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    HEADER_FILL = PatternFill(
        start_color="1E252D", end_color="1E252D", fill_type="solid"
    )
    HEADER_FONT = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    HEADER_BORDER = Border(
        left=Side(style="thin", color="334155"),
        right=Side(style="thin", color="334155"),
        top=Side(style="thin", color="334155"),
        bottom=Side(style="thin", color="334155"),
    )
    DATA_FONT = Font(name="Segoe UI", size=10, color="0F172A")
    LINK_FONT = Font(name="Segoe UI", size=10, color="0969DA", underline="single")
    ZEBRA_EVEN_FILL = PatternFill(
        start_color="F8FAFC", end_color="F8FAFC", fill_type="solid"
    )
    ZEBRA_ODD_FILL = PatternFill(
        start_color="FFFFFF", end_color="FFFFFF", fill_type="solid"
    )
    THIN_BORDER = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    OVERVIEW_SECTION_FILL = PatternFill(
        start_color="14181C", end_color="14181C", fill_type="solid"
    )
    OVERVIEW_SECTION_FONT = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    OVERVIEW_LABEL_FILL = PatternFill(
        start_color="F1F5F9", end_color="F1F5F9", fill_type="solid"
    )
    OVERVIEW_LABEL_FONT = Font(name="Segoe UI", size=10, bold=True, color="475569")

    # Conditional formatting rules
    YES_RULE = CellIsRule(
        operator="equal",
        formula=['"Yes"'],
        stopIfTrue=True,
        fill=PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid"),
        font=Font(name="Segoe UI", size=10, color="15803D", bold=True),
    )
    SPOILER_RULE = CellIsRule(
        operator="equal",
        formula=['"Yes"'],
        stopIfTrue=True,
        fill=PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid"),
        font=Font(name="Segoe UI", size=10, color="B91C1C", bold=True),
    )
    RATING_SCALE_RULE = ColorScaleRule(
        start_type="num",
        start_value=1.0,
        start_color="FED7AA",
        mid_type="num",
        mid_value=3.0,
        mid_color="FEF08A",
        end_type="num",
        end_value=5.0,
        end_color="86EFAC",
    )
    RUNTIME_SCALE_RULE = ColorScaleRule(
        start_type="min",
        start_color="FFFFFF",
        end_type="max",
        end_color="BAE6FD",
    )
    COUNT_DATABAR_RULE = DataBarRule(
        start_type="num",
        start_value=0,
        end_type="max",
        color="00E054",
    )
except ImportError:
    openpyxl = None


class UserDataExcelExporter:
    """Exports comprehensive Letterboxd user data to an Excel (.xlsx) workbook."""

    SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
    EXPORTS_DIR = build_path(SCRIPT_DIR, "exports")
    USERS_FOLDER = build_path(EXPORTS_DIR, "users")

    METHODS = [
        user.User.get_activity,
        user.User.get_activity_following,
        user.User.get_diary,
        user.User.get_wrapped,
        user.User.get_films,
        user.User.get_films_not_rated,
        user.User.get_genre_info,
        user.User.get_liked_films,
        user.User.get_liked_reviews,
        user.User.get_liked_lists,
        user.User.get_lists,
        user.User.get_following,
        user.User.get_followers,
        user.User.get_reviews,
        user.User.get_user_tags,
        user.User.get_watchlist,
    ]

    SOURCE_KEYS = [
        "user",
        "films",
        "diary",
        "reviews",
        "watchlist",
        "liked_films",
        "liked_reviews",
        "liked_lists",
        "lists",
        "following",
        "followers",
    ]

    def __init__(
        self,
        username: str,
        force: bool = False,
        open_file: bool = False,
        fetch_runtime: bool = False,
    ):
        self.username = username
        self.force = force
        self.open_file = open_file
        self.fetch_runtime = fetch_runtime

        self.user_instance = user.User(username)
        self.user_folder = build_path(self.USERS_FOLDER, self.user_instance.username)
        self.sources: dict = {}
        self.wb = None

    def prepare_directories(self) -> None:
        """Create necessary export directory structure."""
        check_and_create_dirs([self.EXPORTS_DIR, self.USERS_FOLDER, self.user_folder])

    def ensure_data(self) -> dict:
        """Ensure all required endpoints are loaded from cache or fetched."""
        user_data_path = build_path(self.user_folder, "user")
        user_data = None if self.force else JsonFile.load(user_data_path)
        if user_data is None:
            user_data = self.user_instance.jsonify()
            JsonFile.save(user_data_path, user_data)

        self.sources = {"user": user_data}

        missing_methods = []
        for method in self.METHODS:
            method_name = method.__name__
            prefix_stripped = strip_prefix(method_name)
            cached = (
                None
                if self.force
                else JsonFile.load(build_path(self.user_folder, prefix_stripped))
            )
            # Invalidate cached diary if runtime was requested but is missing in cached data
            if cached is not None and self.fetch_runtime and prefix_stripped == "diary":
                entries = cached.get("entries", {})
                if any(
                    e.get("runtime") is None
                    for e in entries.values()
                    if isinstance(e, dict)
                ):
                    cached = None

            if cached is not None:
                self.sources[prefix_stripped] = cached
            else:
                missing_methods.append(method)

        if missing_methods:
            action_msg = "Force-fetching" if self.force else "Fetching"
            print(f"{action_msg} {len(missing_methods)} endpoints from Letterboxd...")
            for no, method in enumerate(missing_methods, 1):
                method_name = method.__name__
                prefix_stripped = strip_prefix(method_name)
                print(
                    f"[{no}/{len(missing_methods)}]: Fetching {method_name}...",
                    end="\r",
                )
                try:
                    kwargs = {}
                    if "fetch_runtime" in inspect.signature(method).parameters:
                        kwargs["fetch_runtime"] = self.fetch_runtime

                    data = method(self.user_instance, **kwargs)
                    self.sources[prefix_stripped] = data
                    JsonFile.save(build_path(self.user_folder, prefix_stripped), data)
                    print(f"[{no}/{len(missing_methods)}]: Fetched {method_name}   ")
                except Exception as e:
                    print(
                        f"[{no}/{len(missing_methods)}]: Error fetching {method_name}: {e}"
                    )

        # Fallback load for any remaining required source keys
        for key in self.SOURCE_KEYS:
            if key not in self.sources:
                loaded = JsonFile.load(build_path(self.user_folder, key))
                self.sources[key] = loaded if loaded is not None else {}

        return self.sources

    CENTER_HEADERS = {
        "Year",
        "Release Year",
        "Date",
        "Rating",
        "Liked",
        "Rewatch",
        "Reviewed",
        "Spoiler",
        "Runtime (min)",
        "Log ID",
        "Letterboxd ID",
        "Account Badge",
    }
    RIGHT_HEADERS = {
        "Count",
        "Film Count",
        "Likes",
        "Comments",
        "Watched Films",
        "Followers",
        "Following",
        "Watched",
        "Rating Count",
    }

    @staticmethod
    def _to_int(val) -> int | str:
        if val is None or val == "":
            return ""
        try:
            return int(val)
        except (ValueError, TypeError):
            return str(val)

    @staticmethod
    def _to_float(val) -> float | str:
        if val is None or val == "":
            return ""
        try:
            return float(val)
        except (ValueError, TypeError):
            return str(val)

    def _style_sheet(
        self,
        ws,
        wrap_cols: list | None = None,
        tab_color: str | None = None,
    ) -> None:
        """Apply styles, zebra striping, freeze panes, auto-filters, alignments, and column widths."""
        ws.views.sheetView[0].showGridLines = True
        if tab_color:
            ws.sheet_properties.tabColor = tab_color

        ws.freeze_panes = "A2"
        if ws.max_row > 1:
            ws.auto_filter.ref = ws.dimensions

        ws.row_dimensions[1].height = 26

        wrap_col_indices = set(wrap_cols or [])
        col_alignments = {}

        for col_idx in range(1, ws.max_column + 1):
            h_cell = ws.cell(row=1, column=col_idx)
            h_val = str(h_cell.value or "").strip()

            if h_val in self.CENTER_HEADERS:
                align = "center"
            elif h_val in self.RIGHT_HEADERS:
                align = "right"
            else:
                align = "left"

            col_alignments[col_idx] = (align, h_val)
            h_cell.fill = HEADER_FILL
            h_cell.font = HEADER_FONT
            h_cell.border = HEADER_BORDER
            h_cell.alignment = Alignment(horizontal=align, vertical="center")

        for row_idx in range(2, ws.max_row + 1):
            if (
                row_idx not in ws.row_dimensions
                or not ws.row_dimensions[row_idx].height
            ):
                ws.row_dimensions[row_idx].height = 21

            row_fill = ZEBRA_EVEN_FILL if row_idx % 2 == 0 else ZEBRA_ODD_FILL

            for col_idx in range(1, ws.max_column + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                align, h_val = col_alignments.get(col_idx, ("left", ""))
                v_align = "top" if col_idx in wrap_col_indices else "center"

                cell.fill = row_fill
                cell.font = DATA_FONT
                cell.border = THIN_BORDER
                cell.alignment = Alignment(
                    horizontal=align,
                    vertical=v_align,
                    wrap_text=(col_idx in wrap_col_indices),
                )

                val = cell.value
                if val is not None and val != "":
                    if isinstance(val, float) and "Rating" in h_val:
                        cell.number_format = "0.0"
                    elif isinstance(val, int) and ("Year" in h_val or "ID" in h_val):
                        cell.number_format = "0"
                    elif isinstance(val, int) and (
                        h_val in self.RIGHT_HEADERS
                        or "Count" in h_val
                        or "Runtime" in h_val
                    ):
                        cell.number_format = "#,##0"
                    elif str(val).startswith("http"):
                        cell.hyperlink = str(val)
                        cell.font = LINK_FONT

        for col in ws.columns:
            col_idx = col[0].column
            col_letter = get_column_letter(col_idx)
            h_val = str(col[0].value or "").strip()

            max_len = 0
            for cell in col:
                val = str(cell.value or "")
                line_len = (
                    max((len(line) for line in val.split("\n")), default=0)
                    if "\n" in val
                    else len(val)
                )
                if line_len > max_len:
                    max_len = line_len

            if "URL" in h_val or "Link" in h_val:
                width = min(max(max_len + 2, 20), 38)
            elif col_idx in wrap_col_indices:
                width = 58
            elif h_val in self.CENTER_HEADERS:
                width = max(len(h_val) + 4, 11)
            elif h_val in self.RIGHT_HEADERS:
                width = max(len(h_val) + 4, 14)
            else:
                width = min(max(max_len + 3, 14), 48)

            ws.column_dimensions[col_letter].width = width

    def _build_overview_sheet(self) -> None:
        """Build Sheet 1: User Profile & Statistics Overview."""
        user_data = self.sources.get("user", {})
        films_data = self.sources.get("films", {})
        diary_data = self.sources.get("diary", {})
        reviews_data = self.sources.get("reviews", {})
        watchlist_data = self.sources.get("watchlist", {})
        lists_data = self.sources.get("lists", {})
        following_data = self.sources.get("following", {})
        followers_data = self.sources.get("followers", {})

        ws = self.wb.create_sheet(title="Overview")
        ws.views.sheetView[0].showGridLines = True
        ws.sheet_properties.tabColor = "14181C"

        stats = user_data.get("stats", {})
        favs = user_data.get("favorites", {})
        fav_list = [
            f"{f.get('name', '')} ({f.get('year', '')})".strip()
            for f in favs.values()
            if isinstance(f, dict) and f.get("name")
        ]

        sections = [
            (
                "USER PROFILE",
                [
                    (
                        "Username",
                        user_data.get("username", self.user_instance.username),
                    ),
                    ("Display Name", user_data.get("display_name", "")),
                    (
                        "Account Badge",
                        str(user_data.get("badge", "member")).upper(),
                    ),
                    ("Profile URL", user_data.get("url", "")),
                    ("Bio", user_data.get("bio", "")),
                    ("Location", user_data.get("location", "")),
                    ("Website", user_data.get("website", "")),
                ],
            ),
            (
                "LIBRARY & ACTIVITY STATISTICS",
                [
                    (
                        "Watched Films",
                        self._to_int(films_data.get("count", stats.get("films", 0))),
                    ),
                    (
                        "Rated Films",
                        self._to_int(films_data.get("rating_count", 0)),
                    ),
                    (
                        "Average Rating",
                        self._to_float(films_data.get("rating_average", 0)),
                    ),
                    (
                        "Liked Films",
                        self._to_int(films_data.get("liked_count", 0)),
                    ),
                    ("Diary Entries", self._to_int(diary_data.get("count", 0))),
                    ("Reviews", len(reviews_data.get("reviews", {}))),
                    (
                        "Watchlist",
                        self._to_int(
                            watchlist_data.get("count", stats.get("watchlist", 0))
                        ),
                    ),
                    ("Lists Created", len(lists_data.get("lists", {}))),
                    (
                        "Following",
                        len(following_data) if isinstance(following_data, dict) else 0,
                    ),
                    (
                        "Followers",
                        len(followers_data) if isinstance(followers_data, dict) else 0,
                    ),
                ],
            ),
            (
                "FAVORITE FILMS",
                [(f"Favorite #{idx}", film) for idx, film in enumerate(fav_list, 1)]
                if fav_list
                else [("Favorites", "None selected")],
            ),
        ]

        current_row = 1
        for section_title, rows in sections:
            ws.merge_cells(
                start_row=current_row,
                start_column=1,
                end_row=current_row,
                end_column=2,
            )
            sec_cell = ws.cell(row=current_row, column=1, value=f"  {section_title}")
            sec_cell.fill = OVERVIEW_SECTION_FILL
            sec_cell.font = OVERVIEW_SECTION_FONT
            sec_cell.alignment = Alignment(vertical="center", horizontal="left")
            ws.row_dimensions[current_row].height = 25
            current_row += 1

            for label, val in rows:
                c1 = ws.cell(row=current_row, column=1, value=label)
                c2 = ws.cell(row=current_row, column=2, value=val)

                c1.font = OVERVIEW_LABEL_FONT
                c1.fill = OVERVIEW_LABEL_FILL
                c1.border = THIN_BORDER
                c1.alignment = Alignment(vertical="center", horizontal="left")

                c2.font = DATA_FONT
                c2.border = THIN_BORDER
                c2.alignment = Alignment(vertical="center", horizontal="left")

                if isinstance(val, int):
                    c2.number_format = "#,##0"
                elif isinstance(val, float):
                    c2.number_format = "0.00"
                elif str(val).startswith("http"):
                    c2.hyperlink = str(val)
                    c2.font = LINK_FONT

                ws.row_dimensions[current_row].height = 21
                current_row += 1

            current_row += 1

        ws.column_dimensions["A"].width = 28
        ws.column_dimensions["B"].width = 54

    def _build_films_sheet(self) -> None:
        """Build Sheet 2: Watched Films."""
        films_data = self.sources.get("films", {})
        ws = self.wb.create_sheet(title="Films")
        ws.append(["Film Name", "Year", "Rating", "Liked", "URL", "Letterboxd ID"])
        for m in films_data.get("movies", {}).values():
            ws.append(
                [
                    m.get("name", ""),
                    self._to_int(m.get("year")),
                    self._to_float(m.get("rating")),
                    "Yes" if m.get("liked") else "No",
                    m.get("url", ""),
                    m.get("id", ""),
                ]
            )
        self._style_sheet(ws, tab_color="00E054")
        if ws.max_row > 1:
            ws.conditional_formatting.add(f"C2:C{ws.max_row}", RATING_SCALE_RULE)
            ws.conditional_formatting.add(f"D2:D{ws.max_row}", YES_RULE)

    def _build_diary_sheet(self) -> None:
        """Build Sheet 3: Diary Entries."""
        diary_data = self.sources.get("diary", {})
        ws = self.wb.create_sheet(title="Diary")
        ws.append(
            [
                "Date",
                "Film Name",
                "Release Year",
                "Rating",
                "Liked",
                "Rewatch",
                "Reviewed",
                "Runtime (min)",
                "Log ID",
            ]
        )
        for log_id, entry in diary_data.get("entries", {}).items():
            actions = entry.get("actions", {})
            date_str = str(entry.get("date", "")).replace("T00:00:00.000000Z", "")
            ws.append(
                [
                    date_str,
                    entry.get("name", ""),
                    self._to_int(entry.get("release")),
                    self._to_float(actions.get("rating")),
                    "Yes" if actions.get("liked") else "No",
                    "Yes" if actions.get("rewatched") else "No",
                    "Yes" if actions.get("reviewed") else "No",
                    self._to_int(entry.get("runtime")),
                    log_id,
                ]
            )
        self._style_sheet(ws, tab_color="00E054")
        if ws.max_row > 1:
            ws.conditional_formatting.add(f"D2:D{ws.max_row}", RATING_SCALE_RULE)
            ws.conditional_formatting.add(f"E2:G{ws.max_row}", YES_RULE)
            ws.conditional_formatting.add(f"H2:H{ws.max_row}", RUNTIME_SCALE_RULE)

    def _build_reviews_sheet(self) -> None:
        """Build Sheet 4: Reviews."""
        reviews_data = self.sources.get("reviews", {})
        ws = self.wb.create_sheet(title="Reviews")
        ws.append(
            [
                "Date",
                "Film Name",
                "Release Year",
                "Rating",
                "Spoiler",
                "Review Content",
                "Review URL",
            ]
        )
        for r_data in reviews_data.get("reviews", {}).values():
            movie = r_data.get("movie", {})
            rev_info = r_data.get("review", {})
            date_str = str(r_data.get("date", "")).replace("T00:00:00.000000Z", "")
            ws.append(
                [
                    date_str,
                    movie.get("name", ""),
                    self._to_int(movie.get("release")),
                    self._to_float(r_data.get("rating")),
                    "Yes" if rev_info.get("spoiler") else "No",
                    rev_info.get("content", ""),
                    r_data.get("link", "") or r_data.get("url", ""),
                ]
            )
        self._style_sheet(ws, wrap_cols=[6], tab_color="FF8000")
        if ws.max_row > 1:
            ws.conditional_formatting.add(f"D2:D{ws.max_row}", RATING_SCALE_RULE)
            ws.conditional_formatting.add(f"E2:E{ws.max_row}", SPOILER_RULE)

    def _build_watchlist_sheet(self) -> None:
        """Build Sheet 5: Watchlist."""
        watchlist_data = self.sources.get("watchlist", {})
        ws = self.wb.create_sheet(title="Watchlist")
        ws.append(["Film Name", "Year", "URL", "Letterboxd ID"])
        watch_data = watchlist_data.get("data", {}) or watchlist_data.get("movies", {})
        if isinstance(watch_data, dict):
            for m in watch_data.values():
                ws.append(
                    [
                        m.get("name", ""),
                        self._to_int(m.get("year")),
                        m.get("url", ""),
                        m.get("slug", "") or m.get("id", ""),
                    ]
                )
        self._style_sheet(ws, tab_color="40BCF4")

    def _build_ratings_breakdown_sheet(self) -> None:
        """Build Sheet 6: Ratings Breakdown."""
        films_data = self.sources.get("films", {})
        ws = self.wb.create_sheet(title="Ratings Breakdown")
        ws.append(["Rating", "Count", "Sample Movies"])
        ratings_order = [5.0, 4.5, 4.0, 3.5, 3.0, 2.5, 2.0, 1.5, 1.0, 0.5]
        ratings_map = {r: [] for r in ratings_order}
        for m in films_data.get("movies", {}).values():
            r = m.get("rating")
            if r is not None and float(r) in ratings_map:
                ratings_map[float(r)].append(m.get("name", ""))
        for r_val in ratings_order:
            movies = ratings_map[r_val]
            sample = ", ".join(movies[:5])
            ws.append([f"{r_val} ★", len(movies), sample])
        self._style_sheet(ws, tab_color="2C3440")
        if ws.max_row > 1:
            ws.conditional_formatting.add(f"B2:B{ws.max_row}", COUNT_DATABAR_RULE)

    def _build_liked_films_sheet(self) -> None:
        """Build Sheet 7: Liked Films."""
        liked_films_data = self.sources.get("liked_films", {})
        ws = self.wb.create_sheet(title="Liked Films")
        ws.append(["Film Name", "Year", "Rating", "URL", "Letterboxd ID"])
        for m in liked_films_data.get("movies", {}).values():
            ws.append(
                [
                    m.get("name", ""),
                    self._to_int(m.get("year")),
                    self._to_float(m.get("rating")),
                    m.get("url", ""),
                    m.get("id", ""),
                ]
            )
        self._style_sheet(ws, tab_color="00E054")
        if ws.max_row > 1:
            ws.conditional_formatting.add(f"C2:C{ws.max_row}", RATING_SCALE_RULE)

    def _build_liked_reviews_sheet(self) -> None:
        """Build Sheet 8: Liked Reviews."""
        liked_reviews_data = self.sources.get("liked_reviews", {})
        ws = self.wb.create_sheet(title="Liked Reviews")
        ws.append(["Film Name", "Author", "Rating", "Review Content", "URL"])
        for r in liked_reviews_data.get("reviews", {}).values():
            user_info = r.get("user", {})
            author = (
                user_info.get("display_name")
                or user_info.get("name")
                or user_info.get("username", "")
                if isinstance(user_info, dict)
                else ""
            )
            ws.append(
                [
                    r.get("movie", {}).get("name", "")
                    if isinstance(r.get("movie"), dict)
                    else "",
                    author,
                    self._to_float(r.get("rating")),
                    r.get("review", {}).get("content", "")
                    if isinstance(r.get("review"), dict)
                    else "",
                    r.get("url", ""),
                ]
            )
        self._style_sheet(ws, wrap_cols=[4], tab_color="FF8000")
        if ws.max_row > 1:
            ws.conditional_formatting.add(f"C2:C{ws.max_row}", RATING_SCALE_RULE)

    def _build_lists_sheet(self) -> None:
        """Build Sheet 9: User's Created Lists."""
        lists_data = self.sources.get("lists", {})
        ws = self.wb.create_sheet(title="Lists")
        ws.append(
            ["List Title", "Film Count", "Likes", "Comments", "URL", "Description"]
        )
        for lst in lists_data.get("lists", {}).values():
            ws.append(
                [
                    lst.get("title", ""),
                    self._to_int(lst.get("count")),
                    self._to_int(lst.get("likes")),
                    self._to_int(lst.get("comments")),
                    lst.get("url", ""),
                    lst.get("description", ""),
                ]
            )
        self._style_sheet(ws, wrap_cols=[6], tab_color="40BCF4")

    def _build_liked_lists_sheet(self) -> None:
        """Build Sheet 10: Liked Lists."""
        liked_lists_data = self.sources.get("liked_lists", {})
        ws = self.wb.create_sheet(title="Liked Lists")
        ws.append(
            ["List Title", "Film Count", "Likes", "Comments", "URL", "Description"]
        )
        for lst in liked_lists_data.get("lists", {}).values():
            ws.append(
                [
                    lst.get("title", ""),
                    self._to_int(lst.get("count")),
                    self._to_int(lst.get("likes")),
                    self._to_int(lst.get("comments")),
                    lst.get("url", ""),
                    lst.get("description", ""),
                ]
            )
        self._style_sheet(ws, wrap_cols=[6], tab_color="40BCF4")

    def _build_following_sheet(self) -> None:
        """Build Sheet 11: Following Network."""
        following_data = self.sources.get("following", {})
        ws = self.wb.create_sheet(title="Following")
        ws.append(
            [
                "Username",
                "Display Name",
                "Watched Films",
                "Followers",
                "Following",
                "URL",
            ]
        )
        if isinstance(following_data, dict):
            for u in following_data.values():
                if isinstance(u, dict):
                    ws.append(
                        [
                            u.get("username", ""),
                            u.get("name", ""),
                            self._to_int(u.get("watched")),
                            self._to_int(u.get("followers")),
                            self._to_int(u.get("following")),
                            u.get("url", ""),
                        ]
                    )
        self._style_sheet(ws, tab_color="64748B")

    def _build_followers_sheet(self) -> None:
        """Build Sheet 12: Followers Network."""
        followers_data = self.sources.get("followers", {})
        ws = self.wb.create_sheet(title="Followers")
        ws.append(
            [
                "Username",
                "Display Name",
                "Watched Films",
                "Followers",
                "Following",
                "URL",
            ]
        )
        if isinstance(followers_data, dict):
            for u in followers_data.values():
                if isinstance(u, dict):
                    ws.append(
                        [
                            u.get("username", ""),
                            u.get("name", ""),
                            self._to_int(u.get("watched")),
                            self._to_int(u.get("followers")),
                            self._to_int(u.get("following")),
                            u.get("url", ""),
                        ]
                    )
        self._style_sheet(ws, tab_color="64748B")

    def export(self) -> str | None:
        """Generate workbook, populate sheets, and save to .xlsx file."""
        if openpyxl is None:
            print("\n[Warning] 'openpyxl' is required for Excel export.")
            print("Install it with: pip install openpyxl")
            return None

        self.wb = openpyxl.Workbook()
        self.wb.remove(self.wb.active)  # Remove default blank sheet

        # Build all 12 sheets
        self._build_overview_sheet()
        self._build_films_sheet()
        self._build_diary_sheet()
        self._build_reviews_sheet()
        self._build_watchlist_sheet()
        self._build_ratings_breakdown_sheet()
        self._build_liked_films_sheet()
        self._build_liked_reviews_sheet()
        self._build_lists_sheet()
        self._build_liked_lists_sheet()
        self._build_following_sheet()
        self._build_followers_sheet()

        xlsx_path = build_path(
            self.user_folder, f"{self.user_instance.username}_data.xlsx"
        )
        while True:
            try:
                self.wb.save(xlsx_path)
                break
            except PermissionError:
                print(
                    f"\n[Warning] Permission denied: Cannot write to '{os.path.basename(xlsx_path)}'."
                )
                print("The file appears to be open in Excel or another program.")
                try:
                    choice = (
                        input(
                            "Close the file and press Enter to retry (or 'f' for a new file): "
                        )
                        .strip()
                        .lower()
                    )
                    if choice == "f":
                        import datetime

                        ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
                        xlsx_path = build_path(
                            self.user_folder,
                            f"{self.user_instance.username}_data_{ts}.xlsx",
                        )
                except (KeyboardInterrupt, EOFError):
                    print("\nExport cancelled.")
                    return None
        return xlsx_path

    def run(self) -> str | None:
        """Execute full preparation, collection, and export lifecycle."""
        self.prepare_directories()

        start_time = time.time()
        mode_parts = []
        if self.force:
            mode_parts.append("forcing fresh fetch")
        if self.fetch_runtime:
            mode_parts.append("fetching runtimes")
        mode_str = f" ({', '.join(mode_parts)})" if mode_parts else ""
        print(f"\nPreparing data for @{self.user_instance.username}{mode_str}...")
        self.ensure_data()

        print("Generating Excel spreadsheet (.xlsx)...")
        xlsx_path = self.export()

        if xlsx_path:
            print("\nExport completed successfully!")
            print(f"\tExcel file: {build_click_url(xlsx_path)}")
            print(f"\tTime elapsed: {time.time() - start_time:.2f} seconds\n")

            if self.open_file:
                abs_path = os.path.abspath(xlsx_path)
                webbrowser.open(f"file:///{abs_path.replace(os.sep, '/')}")
        else:
            print("\nExport failed. Make sure 'openpyxl' is installed.")

        return xlsx_path


def main():
    parser = argparse.ArgumentParser(
        description="Exports comprehensive user data to a styled multi-sheet Excel spreadsheet."
    )
    parser.add_argument("username", nargs="?", help="Letterboxd username (positional)")
    parser.add_argument("--user", "-u", help="Letterboxd username (named)")
    parser.add_argument(
        "--force", "-f", action="store_true", help="Force refresh cached data"
    )
    parser.add_argument(
        "--runtime",
        "-r",
        "--fetch-runtime",
        action="store_true",
        dest="fetch_runtime",
        help="Fetch film runtimes for diary entries (requires extra network requests)",
    )
    parser.add_argument(
        "--open",
        "-o",
        action="store_true",
        help="Automatically open the generated Excel file when done",
    )
    args = parser.parse_args()

    if not args_exists():
        print(
            f"Quick usage: python {sys.argv[0]} <username> [--force] [--runtime] [--open]\n"
        )

    username = args.user or args.username or get_input("Enter username: ")
    exporter = UserDataExcelExporter(
        username=username,
        force=args.force,
        open_file=args.open,
        fetch_runtime=args.fetch_runtime,
    )
    exporter.run()

    if not args_exists():
        os.system("pause")


if __name__ == "__main__":
    main()
